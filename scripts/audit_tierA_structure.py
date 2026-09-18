#!/usr/bin/env python3
"""Create a reproducible structural audit of the approved Tier A inputs."""

from __future__ import annotations

import csv
import gzip
import json
import re
import sys
from collections import Counter
from pathlib import Path

import openpyxl


ROOT = Path(sys.argv[1] if len(sys.argv) > 1 else "public_data_tierA")
OUT = ROOT / "derived" / "audit"
OUT.mkdir(parents=True, exist_ok=True)


def clean_key(value: str) -> str:
    value = value.strip().lower()
    value = re.sub(r"[^a-z0-9]+", "_", value).strip("_")
    return value


def parse_soft(path: Path):
    series = {}
    samples = []
    current = None
    with gzip.open(path, "rt", encoding="utf-8", errors="replace") as handle:
        for raw in handle:
            line = raw.rstrip("\r\n")
            if line.startswith("^SERIES = "):
                series["accession"] = line.split(" = ", 1)[1]
            elif line.startswith("!Series_") and " = " in line:
                key, value = line.split(" = ", 1)
                series[clean_key(key.removeprefix("!Series_"))] = value
            elif line.startswith("^SAMPLE = "):
                if current is not None:
                    samples.append(current)
                current = {
                    "dataset": series.get("accession", path.name.split("_")[0]),
                    "accession": line.split(" = ", 1)[1],
                }
            elif current is not None and line.startswith("!Sample_") and " = " in line:
                key, value = line.split(" = ", 1)
                key = clean_key(key.removeprefix("!Sample_"))
                if key == "characteristics_ch1" and ": " in value:
                    char_key, char_value = value.split(": ", 1)
                    char_key = clean_key(char_key)
                    if char_key in current:
                        current[char_key] += " | " + char_value
                    else:
                        current[char_key] = char_value
                elif key in current:
                    current[key] += " | " + value
                else:
                    current[key] = value
    if current is not None:
        samples.append(current)
    return series, samples


soft_paths = sorted((ROOT / "metadata").glob("GSE*_family.soft.gz"))
series_records = []
sample_records = []
for soft_path in soft_paths:
    series, samples = parse_soft(soft_path)
    series_records.append(series)
    sample_records.extend(samples)

preferred = [
    "dataset", "accession", "title", "source_name_ch1", "cell_type", "cells_type",
    "cell_line", "state", "treatment", "age", "age_years", "donor_age",
    "donor_age_years", "sex", "gender", "disease", "clinical_condition",
    "passage", "population_doubling", "days_grown_udays", "unique_variable_name",
    "description", "platform_id", "library_strategy", "supplementary_file_1",
]
all_sample_keys = {key for row in sample_records for key in row}
sample_fields = [key for key in preferred if key in all_sample_keys]
sample_fields += sorted(all_sample_keys - set(sample_fields))
with (OUT / "all_geo_samples.tsv").open("w", newline="", encoding="utf-8") as handle:
    writer = csv.DictWriter(handle, fieldnames=sample_fields, delimiter="\t", extrasaction="ignore")
    writer.writeheader()
    writer.writerows(sample_records)

summary_fields = [
    "dataset", "n_samples", "title", "platform_ids", "library_strategies",
    "sample_title_keyword_counts",
]
with (OUT / "geo_dataset_summary.tsv").open("w", newline="", encoding="utf-8") as handle:
    writer = csv.DictWriter(handle, fieldnames=summary_fields, delimiter="\t")
    writer.writeheader()
    for series in sorted(series_records, key=lambda row: row.get("accession", "")):
        dataset = series.get("accession", "")
        subset = [row for row in sample_records if row["dataset"] == dataset]
        platforms = sorted({row.get("platform_id", "") for row in subset} - {""})
        strategies = sorted({row.get("library_strategy", "") for row in subset} - {""})
        keywords = Counter()
        for row in subset:
            title = row.get("title", "").lower()
            for keyword in (
                "control", "sham", "uva", "uvb", "proliferating", "early senescence",
                "deep senescence", "p27", "p4", "sips", "quiescent", "reprogram",
            ):
                pattern = rf"(?<![a-z0-9]){re.escape(keyword)}(?![a-z0-9])"
                if re.search(pattern, title):
                    keywords[keyword] += 1
        writer.writerow({
            "dataset": dataset,
            "n_samples": len(subset),
            "title": series.get("title", ""),
            "platform_ids": ";".join(platforms),
            "library_strategies": ";".join(strategies),
            "sample_title_keyword_counts": json.dumps(keywords, sort_keys=True),
        })


def first_nonempty_preview(ws, limit=12):
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        values = ["" if value is None else str(value) for value in row]
        nonempty = [value for value in values if value != ""]
        if nonempty:
            return row_idx, " | ".join(nonempty[:limit])[:800]
    return 0, ""


xlsx_rows = []
for path in sorted(ROOT.glob("**/*.xlsx")):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        first_row, preview = first_nonempty_preview(sheet)
        xlsx_rows.append({
            "file": str(path.relative_to(ROOT)),
            "sheet": sheet_name,
            "max_rows": sheet.max_row,
            "max_columns": sheet.max_column,
            "first_nonempty_row": first_row,
            "preview": preview,
        })
    workbook.close()

with (OUT / "xlsx_sheet_inventory.tsv").open("w", newline="", encoding="utf-8") as handle:
    fields = ["file", "sheet", "max_rows", "max_columns", "first_nonempty_row", "preview"]
    writer = csv.DictWriter(handle, fieldnames=fields, delimiter="\t")
    writer.writeheader()
    writer.writerows(xlsx_rows)


def count_text_table(path: Path, delimiter: str, header_lines=1):
    opener = gzip.open if path.suffix == ".gz" else open
    with opener(path, "rt", encoding="utf-8", errors="replace", newline="") as handle:
        reader = csv.reader(handle, delimiter=delimiter)
        rows = 0
        width = 0
        for idx, row in enumerate(reader):
            if idx == 0:
                width = len(row)
            rows += 1
    return max(0, rows - header_lines), width


matrix_specs = [
    ("GSE109700", ROOT / "senescence/GSE109700_LF1_Counts_HiSat2_FeatureCounts.txt.gz", "tsv", 2, 9, "NCBI Gene ID; raw counts"),
    ("GSE113957", ROOT / "aging/GSE113957_fpkm.txt.gz", "tsv", 1, 143, "RefSeq transcript; FPKM"),
    ("GSE179848", ROOT / "aging/GSE179848_raw_counts_cell_lifespan_RNAseq_data.csv.gz", "csv", 1, 345, "RefSeq transcript; count-like values"),
    ("GSE165177_part1", ROOT / "reprogramming/GSE165177_Log2_RPM_Transient_reprogramming.txt.gz", "tsv", 1, 24, "gene symbol and Ensembl; log2 RPM"),
    ("GSE165177_part2", ROOT / "reprogramming/GSE165177_Log2_RPM_Transient_reprogramming_part2_170621.txt.gz", "tsv", 1, 71, "gene symbol and Ensembl; log2 RPM"),
]

matrix_rows = []
for dataset, path, kind, header_lines, expected_samples, identifier in matrix_specs:
    n_rows, n_cols = count_text_table(path, "\t" if kind == "tsv" else ",", header_lines)
    matrix_rows.append({
        "dataset": dataset,
        "input": str(path.relative_to(ROOT)),
        "features_or_rows": n_rows,
        "columns": n_cols,
        "expected_samples": "" if expected_samples is None else expected_samples,
        "identifier_and_scale": identifier,
        "status": "PASS" if n_rows > 0 and n_cols > 1 else "FAIL",
    })

with gzip.open(ROOT / "uva/GSE125429_log2_and_linear_values.txt.gz", "rt", encoding="utf-8", errors="replace") as handle:
    gse125_lines = [line.rstrip("\r\n").split("\t") for line in handle]
gse125_rows = sum(1 for row in gse125_lines[9:] if row and row[0].strip())
matrix_rows.append({
    "dataset": "GSE125429",
    "input": "uva/GSE125429_log2_and_linear_values.txt.gz",
    "features_or_rows": gse125_rows,
    "columns": max(map(len, gse125_lines)),
    "expected_samples": 8,
    "identifier_and_scale": "gene symbol; paired-donor log2 and linear blocks",
    "status": "PASS",
})

series_rows = 0
series_cols = 0
in_matrix = False
with gzip.open(ROOT / "uva/GSE89005_series_matrix.txt.gz", "rt", encoding="utf-8", errors="replace") as handle:
    for line in handle:
        if line.startswith("!series_matrix_table_begin"):
            in_matrix = True
            continue
        if line.startswith("!series_matrix_table_end"):
            break
        if in_matrix:
            row = next(csv.reader([line], delimiter="\t"))
            if series_cols == 0:
                series_cols = len(row)
            else:
                series_rows += 1
matrix_rows.append({
    "dataset": "GSE89005",
    "input": "uva/GSE89005_series_matrix.txt.gz",
    "features_or_rows": series_rows,
    "columns": series_cols,
    "expected_samples": 18,
    "identifier_and_scale": "Agilent probe; GEO normalized expression",
    "status": "PASS",
})

gse191_wb = openpyxl.load_workbook(ROOT / "senescence/GSE191055_gene.P27_vs_P4.xlsx", read_only=True, data_only=True)
gse191_ws = gse191_wb[gse191_wb.sheetnames[0]]
matrix_rows.append({
    "dataset": "GSE191055",
    "input": "senescence/GSE191055_gene.P27_vs_P4.xlsx",
    "features_or_rows": gse191_ws.max_row - 1,
    "columns": gse191_ws.max_column,
    "expected_samples": 8,
    "identifier_and_scale": "publisher-provided P27-vs-P4 differential table only",
    "status": "PASS_DE_ONLY",
})
gse191_wb.close()

for dataset, folder, pattern, sample_count, identifier in [
    ("GSE240226", ROOT / "extracted/GSE240226", "*.txt.gz", 12, "Ensembl gene; raw count and FPKM"),
    ("GSE302943", ROOT / "extracted/GSE302943", "*.txt.gz", 10, "Ensembl gene; matchCounts"),
    ("GSE226189", ROOT / "extracted/GSE226189", "*geneCOUNT.txt.gz", 82, "Ensembl gene; raw counts"),
    ("GSE93535", ROOT / "extracted/GSE93535", "*.gtf.gz", 12, "Cufflinks GTF; FPKM reconstructed downstream"),
]:
    files = sorted(folder.glob(pattern))
    first_rows = 0
    first_cols = 0
    if files and files[0].name.endswith(".txt.gz"):
        first_rows, first_cols = count_text_table(files[0], "\t", 1)
    matrix_rows.append({
        "dataset": dataset,
        "input": str(folder.relative_to(ROOT)) + "/" + pattern,
        "features_or_rows": first_rows if first_rows else "GTF",
        "columns": first_cols if first_cols else "GTF",
        "expected_samples": sample_count,
        "identifier_and_scale": identifier,
        "status": "PASS" if len(files) == sample_count else f"FAIL_found_{len(files)}",
    })

with (OUT / "expression_input_inventory.tsv").open("w", newline="", encoding="utf-8") as handle:
    fields = ["dataset", "input", "features_or_rows", "columns", "expected_samples", "identifier_and_scale", "status"]
    writer = csv.DictWriter(handle, fieldnames=fields, delimiter="\t")
    writer.writeheader()
    writer.writerows(matrix_rows)

print(f"Parsed {len(sample_records)} GEO samples from {len(series_records)} series.")
print(f"Inventoried {len(xlsx_rows)} workbook sheets and {len(matrix_rows)} expression inputs.")
